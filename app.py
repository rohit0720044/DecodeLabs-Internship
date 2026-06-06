from __future__ import annotations

import json
import os
import socket
import sys
import threading
import webbrowser
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import parse_qs, urlparse

try:
    import openpyxl
except ImportError:
    print("Missing dependency: openpyxl")
    print("Install it with: pip install openpyxl")
    raise


ASSIGNMENTS = [
    {
        "id": "assignment-1",
        "title": "Assignment 1",
        "subtitle": "DecodeLabs Assignment 1",
        "file_name": "DecodeLabs Assignment 1.xlsx",
        "accent": "#0f766e",
    },
    {
        "id": "assignment-2",
        "title": "Assignment 2",
        "subtitle": "DecodeLabs Assignment 2",
        "file_name": "DecodeLabs Assignment 2.xlsx",
        "accent": "#7c3aed",
    },
]

MAX_ROWS_PER_SHEET = 250
MAX_COLS_PER_SHEET = 30
APP_DIR = Path(__file__).resolve().parent


def assignment_path(assignment):
    local_path = APP_DIR / assignment["file_name"]
    if local_path.exists():
        return local_path
    return Path.home() / "Desktop" / assignment["file_name"]


def excel_value(value):
    if value is None:
        return ""
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return value


def load_workbook_payload(assignment_id: str):
    assignment = next((item for item in ASSIGNMENTS if item["id"] == assignment_id), None)
    if assignment is None:
        return {"error": "Assignment not found."}, 404

    path = assignment_path(assignment)
    if not path.exists():
        return {"error": f"Excel file not found: {path}"}, 404

    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    sheets = []

    for worksheet in workbook.worksheets:
        max_row = min(worksheet.max_row or 0, MAX_ROWS_PER_SHEET)
        max_col = min(worksheet.max_column or 0, MAX_COLS_PER_SHEET)
        rows = []

        for row in worksheet.iter_rows(
            min_row=1,
            max_row=max_row,
            min_col=1,
            max_col=max_col,
            values_only=True,
        ):
            rows.append([excel_value(cell) for cell in row])

        sheets.append(
            {
                "name": worksheet.title,
                "rows": rows,
                "totalRows": worksheet.max_row or 0,
                "totalCols": worksheet.max_column or 0,
                "shownRows": max_row,
                "shownCols": max_col,
            }
        )

    return {
        "assignment": {
            "id": assignment["id"],
            "title": assignment["title"],
            "subtitle": assignment["subtitle"],
            "accent": assignment["accent"],
            "fileName": path.name,
        },
        "sheets": sheets,
    }, 200


def manifest_payload():
    projects = []
    for item in ASSIGNMENTS:
        path = assignment_path(item)
        projects.append(
            {
                "id": item["id"],
                "title": item["title"],
                "subtitle": item["subtitle"],
                "fileName": item["file_name"],
                "exists": path.exists(),
                "sizeKb": round(path.stat().st_size / 1024, 1) if path.exists() else 0,
                "accent": item["accent"],
            }
        )
    return {"projects": projects}


HTML = r"""<!doctype html>
<html lang="en">
<head>
  <meta charset="utf-8" />
  <meta name="viewport" content="width=device-width, initial-scale=1" />
  <title>DecodeLabs Assignment Viewer</title>
  <style>
    :root {
      color-scheme: light;
      --bg: #f6f7fb;
      --panel: #ffffff;
      --ink: #172033;
      --muted: #667085;
      --line: #d9dee8;
      --soft: #eef2f6;
      --accent: #0f766e;
      --shadow: 0 20px 60px rgba(23, 32, 51, 0.14);
    }

    * { box-sizing: border-box; }

    body {
      margin: 0;
      min-height: 100vh;
      font-family: Inter, ui-sans-serif, system-ui, -apple-system, BlinkMacSystemFont, "Segoe UI", sans-serif;
      color: var(--ink);
      background:
        linear-gradient(135deg, rgba(15, 118, 110, 0.12), rgba(124, 58, 237, 0.08)),
        var(--bg);
    }

    button, input { font: inherit; }

    .app-shell {
      min-height: 100vh;
      display: grid;
      grid-template-rows: auto 1fr;
    }

    header {
      display: flex;
      align-items: center;
      justify-content: space-between;
      gap: 16px;
      padding: 18px clamp(18px, 4vw, 44px);
      border-bottom: 1px solid rgba(217, 222, 232, 0.75);
      background: rgba(255, 255, 255, 0.82);
      backdrop-filter: blur(14px);
      position: sticky;
      top: 0;
      z-index: 5;
    }

    .brand {
      display: flex;
      align-items: center;
      gap: 12px;
      min-width: 0;
    }

    .mark {
      width: 40px;
      height: 40px;
      border-radius: 8px;
      display: grid;
      place-items: center;
      background: var(--accent);
      color: white;
      font-weight: 800;
      letter-spacing: 0;
      flex: 0 0 auto;
    }

    h1 {
      margin: 0;
      font-size: clamp(20px, 2.4vw, 30px);
      line-height: 1.05;
    }

    .subtitle {
      margin-top: 4px;
      color: var(--muted);
      font-size: 13px;
    }

    .change-btn, .project-card, .tab, .pager button {
      border: 1px solid var(--line);
      background: white;
      color: var(--ink);
      cursor: pointer;
      border-radius: 8px;
      transition: transform 0.16s ease, border-color 0.16s ease, box-shadow 0.16s ease;
    }

    .change-btn {
      padding: 10px 14px;
      white-space: nowrap;
    }

    .change-btn:hover, .project-card:hover, .tab:hover, .pager button:hover {
      border-color: color-mix(in srgb, var(--accent) 55%, var(--line));
      box-shadow: 0 10px 24px rgba(23, 32, 51, 0.08);
      transform: translateY(-1px);
    }

    main {
      padding: clamp(18px, 4vw, 44px);
      display: grid;
      gap: 18px;
    }

    .stats {
      display: grid;
      grid-template-columns: repeat(4, minmax(0, 1fr));
      gap: 12px;
    }

    .stat {
      background: rgba(255, 255, 255, 0.78);
      border: 1px solid rgba(217, 222, 232, 0.8);
      border-radius: 8px;
      padding: 14px;
    }

    .stat span {
      display: block;
      color: var(--muted);
      font-size: 12px;
      margin-bottom: 6px;
    }

    .stat strong {
      font-size: 22px;
      line-height: 1;
    }

    .workspace {
      background: var(--panel);
      border: 1px solid var(--line);
      border-radius: 8px;
      overflow: hidden;
      box-shadow: var(--shadow);
      min-height: 520px;
      display: grid;
      grid-template-rows: auto auto 1fr auto;
    }

    .toolbar {
      display: flex;
      align-items: center;
      justify-content: space-between;
      gap: 12px;
      padding: 14px;
      border-bottom: 1px solid var(--line);
      background: #fbfcfe;
    }

    .tabs {
      display: flex;
      gap: 8px;
      overflow-x: auto;
      padding: 12px 14px;
      border-bottom: 1px solid var(--line);
    }

    .tab {
      padding: 9px 12px;
      color: var(--muted);
      white-space: nowrap;
    }

    .tab.active {
      background: color-mix(in srgb, var(--accent) 10%, white);
      border-color: color-mix(in srgb, var(--accent) 45%, var(--line));
      color: var(--accent);
      font-weight: 700;
    }

    .search {
      width: min(360px, 100%);
      border: 1px solid var(--line);
      border-radius: 8px;
      padding: 10px 12px;
      background: white;
    }

    .sheet-meta {
      color: var(--muted);
      font-size: 13px;
    }

    .table-wrap {
      overflow: auto;
      min-height: 0;
    }

    table {
      border-collapse: separate;
      border-spacing: 0;
      width: 100%;
      min-width: 860px;
      font-size: 13px;
    }

    th, td {
      padding: 10px 12px;
      border-right: 1px solid var(--line);
      border-bottom: 1px solid var(--line);
      max-width: 260px;
      overflow: hidden;
      text-overflow: ellipsis;
      white-space: nowrap;
      background: white;
    }

    th {
      position: sticky;
      top: 0;
      z-index: 1;
      background: #f1f5f9;
      text-align: left;
      font-weight: 750;
    }

    tr:hover td { background: #fbfdff; }

    .pager {
      display: flex;
      align-items: center;
      justify-content: space-between;
      gap: 12px;
      padding: 12px 14px;
      border-top: 1px solid var(--line);
      background: #fbfcfe;
      color: var(--muted);
      font-size: 13px;
    }

    .pager-controls {
      display: flex;
      gap: 8px;
    }

    .pager button {
      padding: 8px 10px;
    }

    .empty {
      padding: 36px;
      color: var(--muted);
      text-align: center;
    }

    .modal-backdrop {
      position: fixed;
      inset: 0;
      z-index: 20;
      display: grid;
      place-items: center;
      padding: 20px;
      background: rgba(17, 24, 39, 0.48);
      backdrop-filter: blur(8px);
    }

    .modal {
      width: min(760px, 100%);
      background: white;
      border-radius: 8px;
      box-shadow: var(--shadow);
      overflow: hidden;
    }

    .modal-head {
      padding: 24px;
      border-bottom: 1px solid var(--line);
    }

    .modal-head h2 {
      margin: 0 0 8px;
      font-size: clamp(24px, 4vw, 36px);
    }

    .modal-head p {
      margin: 0;
      color: var(--muted);
    }

    .project-grid {
      display: grid;
      grid-template-columns: repeat(2, minmax(0, 1fr));
      gap: 14px;
      padding: 18px;
    }

    .project-card {
      text-align: left;
      padding: 18px;
      min-height: 170px;
      display: grid;
      align-content: space-between;
    }

    .project-card h3 {
      margin: 0;
      font-size: 22px;
    }

    .project-card p {
      margin: 8px 0 0;
      color: var(--muted);
    }

    .badge {
      display: inline-flex;
      width: fit-content;
      align-items: center;
      gap: 7px;
      padding: 7px 9px;
      border-radius: 8px;
      background: var(--soft);
      color: var(--muted);
      font-size: 12px;
    }

    .badge-dot {
      width: 8px;
      height: 8px;
      border-radius: 999px;
      background: var(--accent);
    }

    .loading {
      padding: 44px;
      text-align: center;
      color: var(--muted);
    }

    @media (max-width: 760px) {
      header, .toolbar, .pager {
        align-items: stretch;
        flex-direction: column;
      }

      .stats, .project-grid {
        grid-template-columns: 1fr;
      }

      .change-btn, .search {
        width: 100%;
      }
    }
  </style>
</head>
<body>
  <div class="app-shell">
    <header>
      <div class="brand">
        <div class="mark">DL</div>
        <div>
          <h1 id="pageTitle">DecodeLabs Assignments</h1>
          <div class="subtitle" id="pageSubtitle">Choose an assignment to open</div>
        </div>
      </div>
      <button class="change-btn" id="chooseBtn">Change assignment</button>
    </header>

    <main>
      <section class="stats" id="stats"></section>
      <section class="workspace">
        <div class="toolbar">
          <div>
            <strong id="activeSheetTitle">No sheet selected</strong>
            <div class="sheet-meta" id="sheetMeta">Open a project to view workbook data.</div>
          </div>
          <input class="search" id="searchInput" placeholder="Search current sheet..." />
        </div>
        <div class="tabs" id="tabs"></div>
        <div class="table-wrap" id="tableWrap">
          <div class="loading">Waiting for assignment selection...</div>
        </div>
        <div class="pager">
          <span id="pageInfo">Page 1</span>
          <div class="pager-controls">
            <button id="prevBtn">Previous</button>
            <button id="nextBtn">Next</button>
          </div>
        </div>
      </section>
    </main>
  </div>

  <div class="modal-backdrop" id="modal">
    <div class="modal">
      <div class="modal-head">
        <h2>Select Project</h2>
        <p>Open Assignment 1 or Assignment 2 from your DecodeLabs Excel files.</p>
      </div>
      <div class="project-grid" id="projectGrid"></div>
    </div>
  </div>

  <script>
    const state = {
      manifest: null,
      workbook: null,
      activeSheetIndex: 0,
      page: 1,
      pageSize: 25,
      query: ""
    };

    const els = {
      modal: document.getElementById("modal"),
      projectGrid: document.getElementById("projectGrid"),
      chooseBtn: document.getElementById("chooseBtn"),
      pageTitle: document.getElementById("pageTitle"),
      pageSubtitle: document.getElementById("pageSubtitle"),
      stats: document.getElementById("stats"),
      tabs: document.getElementById("tabs"),
      tableWrap: document.getElementById("tableWrap"),
      activeSheetTitle: document.getElementById("activeSheetTitle"),
      sheetMeta: document.getElementById("sheetMeta"),
      searchInput: document.getElementById("searchInput"),
      pageInfo: document.getElementById("pageInfo"),
      prevBtn: document.getElementById("prevBtn"),
      nextBtn: document.getElementById("nextBtn")
    };

    async function init() {
      const response = await fetch("/api/projects");
      state.manifest = await response.json();
      renderProjectCards();
    }

    function renderProjectCards() {
      els.projectGrid.innerHTML = state.manifest.projects.map(project => `
        <button class="project-card" style="--accent:${project.accent}" data-id="${project.id}" ${project.exists ? "" : "disabled"}>
          <span class="badge"><span class="badge-dot"></span>${project.exists ? `${project.sizeKb} KB` : "File missing"}</span>
          <span>
            <h3>${escapeHtml(project.title)}</h3>
            <p>${escapeHtml(project.fileName)}</p>
          </span>
        </button>
      `).join("");

      els.projectGrid.querySelectorAll(".project-card").forEach(card => {
        card.addEventListener("click", () => loadAssignment(card.dataset.id));
      });
    }

    async function loadAssignment(id) {
      els.modal.style.display = "none";
      els.tableWrap.innerHTML = `<div class="loading">Loading workbook...</div>`;

      const response = await fetch(`/api/workbook?id=${encodeURIComponent(id)}`);
      const payload = await response.json();

      if (!response.ok) {
        els.tableWrap.innerHTML = `<div class="empty">${escapeHtml(payload.error || "Could not load workbook.")}</div>`;
        return;
      }

      state.workbook = payload;
      state.activeSheetIndex = 0;
      state.page = 1;
      state.query = "";
      els.searchInput.value = "";
      document.documentElement.style.setProperty("--accent", payload.assignment.accent);
      els.pageTitle.textContent = payload.assignment.title;
      els.pageSubtitle.textContent = payload.assignment.fileName;
      renderAll();
    }

    function renderAll() {
      renderStats();
      renderTabs();
      renderTable();
    }

    function renderStats() {
      const sheets = state.workbook.sheets;
      const rowTotal = sheets.reduce((sum, sheet) => sum + sheet.totalRows, 0);
      const colTotal = sheets.reduce((sum, sheet) => sum + sheet.totalCols, 0);
      const biggest = sheets.reduce((best, sheet) => sheet.totalRows > best.totalRows ? sheet : best, sheets[0]);

      const stats = [
        ["Workbook", state.workbook.assignment.subtitle],
        ["Sheets", sheets.length],
        ["Total Rows", rowTotal.toLocaleString()],
        ["Largest Sheet", biggest ? biggest.name : "-"]
      ];

      els.stats.innerHTML = stats.map(([label, value]) => `
        <div class="stat"><span>${escapeHtml(label)}</span><strong>${escapeHtml(String(value))}</strong></div>
      `).join("");
    }

    function renderTabs() {
      els.tabs.innerHTML = state.workbook.sheets.map((sheet, index) => `
        <button class="tab ${index === state.activeSheetIndex ? "active" : ""}" data-index="${index}">
          ${escapeHtml(sheet.name)}
        </button>
      `).join("");

      els.tabs.querySelectorAll(".tab").forEach(tab => {
        tab.addEventListener("click", () => {
          state.activeSheetIndex = Number(tab.dataset.index);
          state.page = 1;
          renderAll();
        });
      });
    }

    function getFilteredRows(sheet) {
      const rows = (sheet.rows || []).slice(1);
      if (!state.query.trim()) return rows;
      const query = state.query.trim().toLowerCase();
      return rows.filter(row => row.some(cell => String(cell ?? "").toLowerCase().includes(query)));
    }

    function renderTable() {
      const sheet = state.workbook.sheets[state.activeSheetIndex];
      const rows = getFilteredRows(sheet);
      const totalPages = Math.max(1, Math.ceil(rows.length / state.pageSize));
      state.page = Math.min(state.page, totalPages);

      const start = (state.page - 1) * state.pageSize;
      const visibleRows = rows.slice(start, start + state.pageSize);
      const header = sheet.rows[0] || [];

      els.activeSheetTitle.textContent = sheet.name;
      els.sheetMeta.textContent = `${sheet.totalRows.toLocaleString()} rows x ${sheet.totalCols.toLocaleString()} columns. Showing up to ${sheet.shownRows.toLocaleString()} rows and ${sheet.shownCols.toLocaleString()} columns.`;
      els.pageInfo.textContent = `Page ${state.page} of ${totalPages} | ${rows.length.toLocaleString()} visible rows`;
      els.prevBtn.disabled = state.page <= 1;
      els.nextBtn.disabled = state.page >= totalPages;

      if (!visibleRows.length) {
        els.tableWrap.innerHTML = `<div class="empty">No matching rows found.</div>`;
        return;
      }

      els.tableWrap.innerHTML = `
        <table>
          <thead>
            <tr>${header.map((cell, index) => `<th>${escapeHtml(String(cell || `Column ${index + 1}`))}</th>`).join("")}</tr>
          </thead>
          <tbody>
            ${visibleRows.map(row => `
              <tr>${header.map((_, index) => `<td title="${escapeHtml(String(row[index] ?? ""))}">${escapeHtml(String(row[index] ?? ""))}</td>`).join("")}</tr>
            `).join("")}
          </tbody>
        </table>
      `;
    }

    function escapeHtml(value) {
      return value
        .replaceAll("&", "&amp;")
        .replaceAll("<", "&lt;")
        .replaceAll(">", "&gt;")
        .replaceAll('"', "&quot;")
        .replaceAll("'", "&#039;");
    }

    els.chooseBtn.addEventListener("click", () => {
      els.modal.style.display = "grid";
    });

    els.searchInput.addEventListener("input", event => {
      state.query = event.target.value;
      state.page = 1;
      renderTable();
    });

    els.prevBtn.addEventListener("click", () => {
      state.page = Math.max(1, state.page - 1);
      renderTable();
    });

    els.nextBtn.addEventListener("click", () => {
      state.page += 1;
      renderTable();
    });

    init();
  </script>
</body>
</html>
"""


class AppHandler(BaseHTTPRequestHandler):
    def log_message(self, format, *args):
        return

    def send_json(self, payload, status=200):
        body = json.dumps(payload).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def do_GET(self):
        parsed = urlparse(self.path)

        if parsed.path == "/":
            body = HTML.encode("utf-8")
            self.send_response(200)
            self.send_header("Content-Type", "text/html; charset=utf-8")
            self.send_header("Content-Length", str(len(body)))
            self.end_headers()
            self.wfile.write(body)
            return

        if parsed.path == "/api/projects":
            self.send_json(manifest_payload())
            return

        if parsed.path == "/api/workbook":
            assignment_id = parse_qs(parsed.query).get("id", [""])[0]
            payload, status = load_workbook_payload(assignment_id)
            self.send_json(payload, status)
            return

        self.send_response(404)
        self.end_headers()


def find_free_port(start=8000):
    for port in range(start, start + 100):
        with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as sock:
            try:
                sock.bind(("127.0.0.1", port))
                return port
            except OSError:
                continue
    raise RuntimeError("No free local port found.")


def main():
    port = int(os.environ.get("PORT", find_free_port()))
    server = ThreadingHTTPServer(("127.0.0.1", port), AppHandler)
    url = f"http://127.0.0.1:{port}"
    print(f"DecodeLabs Assignment Viewer running at {url}")
    print("Press Ctrl+C to stop the server.")
    threading.Timer(0.8, lambda: webbrowser.open(url)).start()

    try:
        server.serve_forever()
    except KeyboardInterrupt:
        print("\nServer stopped.")
        server.server_close()
        sys.exit(0)


if __name__ == "__main__":
    main()
